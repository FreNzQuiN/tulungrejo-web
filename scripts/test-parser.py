"""Test parser: verify Excel data extraction without DB dependency."""

import openpyxl
import re
import json
from collections import Counter
from datetime import datetime

# ── Config ──
SPOP_FILE = ".secret/ENTRY SPOP LSPOP DESA TULUNGREJO.xlsx"
PBB_FILE = ".secret/PBB-P2 KELURAHAN-WATES.xlsx"

BLOK_TO_DUSUN = {
    "001": "Tulungrejo",
    "002": "Tulungrejo",
    "003": "Tulungrejo",
    "004": "Tulungrejo",
    "005": "Tulungrejo",
    "006": "Sidodadi",
    "007": "Sidodadi",
    "008": "Sidodadi",
    "009": "Sidodadi",
    "010": "Sidodadi",
    "011": "Sidodadi",
    "012": "Sidodadi",
    "013": "TumpakGatho",
}


def sanitize_blok(v):
    if v is None:
        return None
    s = str(v).strip().zfill(3)
    return s if re.match(r"^\d{3}$", s) else None


def sanitize_nobidang(v):
    if v is None:
        return None
    s = str(v).strip().zfill(4)
    return s if re.match(r"^\d{4}$", s) else None


# ── 1. SPOP ──
wb = openpyxl.load_workbook(SPOP_FILE, data_only=True)
ws_spop = wb["SPOP"]
print(f"SPOP rows: {ws_spop.max_row}")

rows = list(ws_spop.iter_rows(min_row=1, max_row=ws_spop.max_row, values_only=True))
# data starts at row index 3 (0-based) = Excel row 4
fields_count = 0
seen = set()
fields_sample = []
for row in rows[3:]:
    if not row:
        continue
    blok = sanitize_blok(row[7])
    nob = sanitize_nobidang(row[8])
    if not blok or not nob:
        continue
    key = f"{blok}|{nob}"
    if key in seen:
        continue
    seen.add(key)
    owner = str(row[13] or "").strip()
    if not owner:
        continue
    fields_count += 1
    if len(fields_sample) < 3:
        addr_parts = [
            str(row[i] or "").strip()
            for i in (14, 15, 16, 17, 18)
            if row[i] is not None
        ]
        owner_addr = ", ".join(a for a in addr_parts if a) or None
        area = round(float(row[22] or 0), 2) if row[22] else None
        fields_sample.append(
            {
                "nop": f"35.05.050.005.{blok}.{nob}",
                "owner": owner,
                "owner_addr": owner_addr,
                "address": str(row[19] or "").strip(),
                "rw": str(row[20] or "").strip() or None,
                "rt": str(row[21] or "").strip() or None,
                "dusun": BLOK_TO_DUSUN.get(blok, "Tulungrejo"),
                "blok": blok,
                "nob": nob,
                "land_area": area,
                "znt": str(row[23] or "").strip() or None,
                "jenis_tanah": int(row[24]) if row[24] is not None else None,
            }
        )

print(f"Fields parsed: {fields_count}")
print("Sample fields:", json.dumps(fields_sample, indent=2, default=str))

# ── 2. LSPOP ──
ws_lspop = wb["LSPOP"]
print(f"\nLSPOP rows: {ws_lspop.max_row}")
lrows = list(ws_lspop.iter_rows(min_row=1, max_row=ws_lspop.max_row, values_only=True))
# data starts at row index 3
lspop_agg = {}  # key -> {area, count}
for row in lrows[3:]:
    if not row:
        continue
    blok = sanitize_blok(row[9])  # col 9 = blok in LSPOP
    nob = sanitize_nobidang(row[10])  # col 10 = no_bidang
    if not blok or not nob:
        continue
    key = f"{blok}|{nob}"
    area_raw = row[16]
    if area_raw is None:
        continue
    area = float(area_raw)
    if area == 0:
        continue
    if key in lspop_agg:
        lspop_agg[key]["area"] += area
        lspop_agg[key]["count"] += 1
    else:
        lspop_agg[key] = {"area": area, "count": 1}

print(f"LSPOP aggregated: {len(lspop_agg)} fields")

# ── 3. Match SPOP + LSPOP ──
match_count = 0
for key in seen:
    if key in lspop_agg:
        match_count += 1
print(f"Fields with building data: {match_count}")

if fields_count > 0:
    blok_list = [key.split("|")[0] for key in seen]
    blok_dist = Counter(blok_list)
    print(f"\nDistribution per blok: {dict(sorted(blok_dist.items()))}")

# ── 4. PBB-P2 ──
print(f"\n=== PBB-P2 ===")
wb2 = openpyxl.load_workbook(PBB_FILE, data_only=True)
ws2 = wb2[wb2.sheetnames[0]]
prows = list(ws2.iter_rows(min_row=1, max_row=ws2.max_row, values_only=True))
print(f"PBB-P2 rows: {ws2.max_row}, cols: {ws2.max_column}")

target_row = None
for row in prows:
    if not row or len(row) < 4:
        continue
    col2 = str(row[2] or "").strip()
    col3 = str(row[3] or "").strip()
    if col2 == "005" or col3.upper() == "TULUNGREJO":
        target_row = row
        break

if target_row:
    print("Found TULUNGREJO row:")
    print(f"  kode_kec:    {target_row[0]}")
    print(f"  kecamatan:   {target_row[1]}")
    print(f"  kode_desa:   {target_row[2]}")
    print(f"  desa:        {target_row[3]}")
    print(f"  PBB:         {target_row[4]}")
    print(f"  BAYAR:       {target_row[5]}")
    print(f"  %:           {target_row[6]}")
    print(f"  KURANG BAYAR:{target_row[7]}")
    print(f"  SPPT:        {target_row[8]}")
    print(f"  DIBAYAR:     {target_row[9]}")
    print(f"  SISA SPPT:   {target_row[10]}")
else:
    print("TULUNGREJO row NOT FOUND")

print("\n✅ Parsing test complete")
